import csv
import io
from datetime import timedelta

from django.http import HttpResponse
from django.utils import timezone

from .models import StockMovement
from .services import SIZE_ORDER, holdings, items_with_balance, status

# column kinds: t=text n=number d=date s=status m=movement type


def stock_report():
    items = sorted(
        items_with_balance(),
        key=lambda i: (i.uniform_type.name, SIZE_ORDER.index(i.size) if i.size in SIZE_ORDER else 99),
    )
    rows = [[i.uniform_type.name, i.size, i.total, i.reorder_level, status(i)] for i in items]
    return {
        "kind": "stock", "title": "Stock report", "sheet": "Stock",
        "subtitle": f"Balances as at {timezone.localdate():%d %b %Y}",
        "columns": [("Uniform type", "t"), ("Size", "t"), ("In stock", "n"), ("Reorder level", "n"), ("Status", "s")],
        "rows": rows,
        "summary": [
            f"{sum(i.total for i in items)} pieces in stock",
            f"{sum(1 for r in rows if r[4] == 'LOW')} low",
            f"{sum(1 for r in rows if r[4] == 'OUT')} out of stock",
        ],
    }


def movement_report(start, end, mtype=""):
    qs = (
        StockMovement.objects.select_related("item__uniform_type", "recorded_by")
        .filter(date__gte=start, date__lte=end)
        .order_by("date", "id")
    )
    if mtype in dict(StockMovement.TYPES):
        qs = qs.filter(movement_type=mtype)
    rows, totals = [], {t: 0 for t, _ in StockMovement.TYPES}
    for m in qs:
        totals[m.movement_type] += m.quantity
        rows.append([
            m.date, m.get_movement_type_display(), m.item.uniform_type.name, m.item.size,
            m.quantity, m.issued_to, m.note, m.recorded_by.get_username() if m.recorded_by else "",
        ])
    return {
        "kind": "movements", "title": "Movement report", "sheet": "Movements",
        "subtitle": f"{start:%d %b %Y} to {end:%d %b %Y}",
        "columns": [("Date", "d"), ("Type", "m"), ("Uniform type", "t"), ("Size", "t"),
                    ("Qty", "n"), ("Person", "t"), ("Note", "t"), ("Recorded by", "t")],
        "rows": rows,
        "summary": [f"{label} {totals[t]}" for t, label in StockMovement.TYPES],
    }


def holdings_report():
    data = holdings()
    rows = [[h["person"], h["type"], h["size"], h["qty"]] for h in data]
    return {
        "kind": "holdings", "title": "Items on Issue", "sheet": "Holdings",
        "subtitle": "Issued minus returned, per person (names are matched ignoring capitals)",
        "columns": [("Person", "t"), ("Uniform type", "t"), ("Size", "t"), ("Pieces held", "n")],
        "rows": rows,
        "summary": [f"{sum(h['qty'] for h in data)} pieces out", f"{len({h['person'].lower() for h in data})} people"],
    }


# ---------- presentation / export ----------------------------------------
GOOD = {"Received", "Returned", "OK"}
BAD = {"Issued", "Written off", "OUT"}


def html_rows(rep):
    out = []
    for row in rep["rows"]:
        cells = []
        for (name, kind), v in zip(rep["columns"], row):
            cls = "r" if kind == "n" else ""
            if kind in ("s", "m"):
                cls = "good" if v in GOOD else "bad" if v in BAD else "warn"
            if kind == "d":
                v = v.strftime("%d %b %Y")
            cells.append({"v": v, "cls": cls})
        out.append(cells)
    return out


def _filename(rep, ext):
    return f"uniform-{rep['kind']}-{timezone.localdate():%Y%m%d}.{ext}"


def to_csv(rep):
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(rep, "csv")}"'
    resp.write("\ufeff")  # so Excel reads it as UTF-8
    w = csv.writer(resp)
    w.writerow([rep["title"], rep["subtitle"]])
    w.writerow([c for c, _ in rep["columns"]])
    w.writerows(rep["rows"])
    return resp


def to_xlsx(rep):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse("Excel export needs openpyxl:  pip install openpyxl", status=501)

    wb = Workbook()
    ws = wb.active
    ws.title = rep["sheet"]
    ws.append([f"Uniform Inventory - {rep['title']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([rep["subtitle"]])
    ws.append([])
    ws.append([c for c, _ in rep["columns"]])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(vertical="center")
    for row in rep["rows"]:
        ws.append(row)
    for idx, (_, kind) in enumerate(rep["columns"], start=1):
        letter = get_column_letter(idx)
        width = max(len(str(ws.cell(r, idx).value or "")) for r in range(4, ws.max_row + 1))
        ws.column_dimensions[letter].width = min(max(width + 3, 10), 45)
        for r in range(5, ws.max_row + 1):
            if kind == "d":
                ws.cell(r, idx).number_format = "dd mmm yyyy"
            if kind == "n":
                ws.cell(r, idx).alignment = Alignment(horizontal="right")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(rep['columns']))}{max(ws.max_row, 4)}"
    ws.append([])
    ws.append(["Summary: " + "  |  ".join(rep["summary"])])

    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{_filename(rep, "xlsx")}"'
    return resp


def parse_date(value, default):
    try:
        return timezone.datetime.fromisoformat(value).date() if value else default
    except (TypeError, ValueError):
        return default


def default_range():
    today = timezone.localdate()
    return today - timedelta(days=29), today
